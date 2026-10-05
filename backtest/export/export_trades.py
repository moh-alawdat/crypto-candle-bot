#!/usr/bin/env python3
"""
Export the Phase 3 backtest (unchanged rules) to an Excel workbook,
mutaz_backtest.xlsx, covering three consecutive 180-day windows:
the last 180 days, the 180 days before that, and the 180 days before that.

Workbook layout:
  * "Summary" sheet: one row per window + a comparison with Gemini's claim.
  * One sheet per window (named by its date range), one row per trade.

Balances follow "Gemini's plan": each window STARTS at $70,000, 10x leverage,
full compounding, fees included -> a win multiplies the balance by 1.09, a loss
by 0.89. (Each window resets to $70,000, per this task's spec. Note this differs
from backtest_windows.py, which compounds continuously across windows; trade
counts, win rates and per-window growth factors are identical either way.)

Run with:
  uv run --with openpyxl --with pandas --with requests export_trades.py
"""

import time
from datetime import datetime, timedelta, timezone

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# --- make repo root + backtest/ subfolders importable (flat layout preserved) ---
import sys as _sys
import pathlib as _pathlib
_ROOT = next(p for p in _pathlib.Path(__file__).resolve().parents
             if (p / "binance_logic.py").exists())
for _d in (_ROOT, _ROOT / "backtest", _ROOT / "backtest" / "runs",
           _ROOT / "backtest" / "export", _ROOT / "backtest" / "checks"):
    if str(_d) not in _sys.path:
        _sys.path.insert(0, str(_d))

from binance_logic import fetch_klines, build_rows
from signals import build_frame, detect_signals, SYMBOL, INTERVAL, OPEN_TIME_FMT
from backtest import simulate, longest_losing_streak, TP_PCT, SL_PCT
from backtest_periods import window_indices
from backtest_windows import TOTAL  # same fetch depth -> identical indicators/trades

WINDOW_DAYS = 180
N_WINDOWS = 3
START_BALANCE = 70000.0
WIN_MULT = 1.09   # +0.9% net * 10x
LOSS_MULT = 0.89  # -1.1% net * 10x
_OUTPUT_DIR = _ROOT / "output"
_OUTPUT_DIR.mkdir(exist_ok=True)
XLSX_PATH = str(_OUTPUT_DIR / "mutaz_backtest.xlsx")

GEMINI_CLAIM = "$70,000 -> $472,500 in 6 months"
GEMINI_END = 472500.0

# Styles
BOLD = Font(bold=True)
HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
WIN_FILL = PatternFill("solid", fgColor="C6EFCE")   # light green
LOSS_FILL = PatternFill("solid", fgColor="FFC7CE")  # light red
CENTER = Alignment(horizontal="center")

# Number formats
PRICE_FMT = "#,##0.00"
MONEY_FMT = '"$"#,##0.00'
PCT_FMT = '0.00"%"'
PCT1_FMT = '0.0"%"'

# (header, width) for the per-window trade sheets
TRADE_COLS = [
    ("Trade #", 9), ("Direction", 11), ("Signal Time (UTC)", 20),
    ("Entry Time (UTC)", 20), ("Entry Price", 14), ("Take-Profit", 14),
    ("Stop-Loss", 14), ("Exit Time (UTC)", 20), ("Exit Price", 14),
    ("Result", 9), ("% After Fees", 13), ("Balance Before", 17),
    ("Balance After", 17),
]

# (header, width) for the Summary sheet
SUMMARY_COLS = [
    ("Dates", 22), ("Trades", 9), ("Wins", 7), ("Losses", 8),
    ("Win Rate", 11), ("% (No Leverage)", 16), ("Starting Balance", 17),
    ("Ending Balance", 17), ("% at 10x", 11), ("Longest Losing Streak", 22),
]


def tp_sl(direction, entry):
    """Take-profit and stop-loss prices for a trade."""
    if direction == "LONG":
        return entry * (1 + TP_PCT / 100), entry * (1 - SL_PCT / 100)
    return entry * (1 - TP_PCT / 100), entry * (1 + SL_PCT / 100)


def load_windows():
    """Fetch once and return a list of windows (newest first), each with trades."""
    raw = fetch_klines(SYMBOL, INTERVAL, TOTAL)
    if raw and raw[-1][6] > int(time.time() * 1000):
        raw = raw[:-1]
    rows = build_rows(raw)
    df = build_frame(rows)
    open_times = [datetime.strptime(t, OPEN_TIME_FMT).replace(tzinfo=timezone.utc)
                  for t in df["open_time"]]
    time_to_idx = {t: i for i, t in enumerate(df["open_time"])}
    now = datetime.now(timezone.utc)

    windows = []
    for k in range(N_WINDOWS):  # k=0 is the last 180 days
        end_time = now - timedelta(days=WINDOW_DAYS * k)
        start_time = now - timedelta(days=WINDOW_DAYS * (k + 1))
        start_idx, end_idx = window_indices(open_times, start_time, end_time)
        signals, _ = detect_signals(df, start_idx, end_idx)
        for s in signals:
            s["_idx"] = time_to_idx[s["open_time (UTC)"]]
        trades, _, _ = simulate(rows, signals)

        sd = df["open_time"].iloc[start_idx][:10]
        ed = df["open_time"].iloc[end_idx - 1][:10]
        name = (f'{datetime.strptime(sd, "%Y-%m-%d").strftime("%b %Y")} - '
                f'{datetime.strptime(ed, "%Y-%m-%d").strftime("%b %Y")}')
        windows.append({"name": name, "start": sd, "end": ed, "trades": trades})
    return windows


def style_header(ws, cols):
    """Write a bold, filled, frozen header row from (name, width) pairs."""
    for c, (name, width) in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font = BOLD
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.freeze_panes = "A2"


def write_trade_sheet(ws, trades):
    """Fill a window sheet with one row per trade; returns ending balance + stats."""
    style_header(ws, TRADE_COLS)

    bal = START_BALANCE
    wins = 0
    for i, t in enumerate(trades, start=1):
        is_win = t["result"] == "TP"
        wins += is_win
        tp, sl = tp_sl(t["direction"], t["entry_price"])
        before = bal
        bal *= WIN_MULT if is_win else LOSS_MULT

        values = [
            i, t["direction"], t["signal_time"], t["entry_time"],
            t["entry_price"], round(tp, 2), round(sl, 2), t["exit_time"],
            t["exit_price"], "WIN" if is_win else "LOSS",
            t["pct_after_fees"], round(before, 2), round(bal, 2),
        ]
        row = i + 1
        for c, v in enumerate(values, start=1):
            ws.cell(row=row, column=c, value=v)

        # Number formats
        for c in (5, 6, 7, 9):
            ws.cell(row=row, column=c).number_format = PRICE_FMT
        ws.cell(row=row, column=11).number_format = PCT_FMT
        ws.cell(row=row, column=12).number_format = MONEY_FMT
        ws.cell(row=row, column=13).number_format = MONEY_FMT

        # Row colouring
        fill = WIN_FILL if is_win else LOSS_FILL
        for c in range(1, len(TRADE_COLS) + 1):
            ws.cell(row=row, column=c).fill = fill

    stats = {
        "trades": len(trades),
        "wins": wins,
        "losses": len(trades) - wins,
        "win_rate": (wins / len(trades) * 100) if trades else 0.0,
        "no_lev": sum(t["pct_after_fees"] for t in trades),
        "end": bal,
        "lev_pct": (bal / START_BALANCE - 1) * 100,
        "streak": longest_losing_streak(trades),
    }
    return stats


def write_summary(ws, windows, stats_by_window):
    style_header(ws, SUMMARY_COLS)

    for i, w in enumerate(windows, start=1):
        s = stats_by_window[w["name"]]
        row = i + 1
        values = [
            w["name"], s["trades"], s["wins"], s["losses"], s["win_rate"],
            s["no_lev"], START_BALANCE, round(s["end"], 2), s["lev_pct"],
            s["streak"],
        ]
        for c, v in enumerate(values, start=1):
            ws.cell(row=row, column=c, value=v)
        ws.cell(row=row, column=5).number_format = PCT1_FMT
        ws.cell(row=row, column=6).number_format = PCT_FMT
        ws.cell(row=row, column=7).number_format = MONEY_FMT
        ws.cell(row=row, column=8).number_format = MONEY_FMT
        ws.cell(row=row, column=9).number_format = PCT_FMT

    # Gemini comparison block, a couple of rows below the table.
    base = len(windows) + 3
    best = max(windows, key=lambda w: stats_by_window[w["name"]]["end"])
    best_stats = stats_by_window[best["name"]]
    gemini_pct = (GEMINI_END / START_BALANCE - 1) * 100

    lines = [
        ("Gemini's claim:", f"{GEMINI_CLAIM}  (+{gemini_pct:.0f}%)"),
        ("Our best 180-day window (10x):",
         f"${best_stats['end']:,.0f}  (+{best_stats['lev_pct']:.1f}%)  [{best['name']}]"),
        ("Gap vs. claim:",
         f"${GEMINI_END - best_stats['end']:,.0f} short of Gemini's ending balance"),
    ]
    for j, (label, text) in enumerate(lines):
        r = base + j
        lc = ws.cell(row=r, column=1, value=label)
        lc.font = BOLD
        ws.cell(row=r, column=2, value=text)


def main():
    windows = load_windows()

    wb = Workbook()
    summary_ws = wb.active
    summary_ws.title = "Summary"

    stats_by_window = {}
    for w in windows:
        ws = wb.create_sheet(title=w["name"])
        stats_by_window[w["name"]] = write_trade_sheet(ws, w["trades"])

    write_summary(summary_ws, windows, stats_by_window)
    wb.save(XLSX_PATH)

    # ---- console summary for cross-checking against backtest_windows.py ----
    print(f"\nSaved {XLSX_PATH} with sheets: "
          f"{['Summary'] + [w['name'] for w in windows]}\n")
    print(f"{'window':<22}{'trades':>7}{'W':>4}{'L':>4}{'win%':>8}"
          f"{'growth x':>10}{'end $ (reset 70k)':>20}")
    for w in windows:
        s = stats_by_window[w["name"]]
        growth = s["end"] / START_BALANCE
        print(f"{w['name']:<22}{s['trades']:>7}{s['wins']:>4}{s['losses']:>4}"
              f"{s['win_rate']:>7.1f}%{growth:>10.5f}{s['end']:>20,.2f}")


if __name__ == "__main__":
    main()
