#!/usr/bin/env python3
"""
Simpler Excel export: mutaz_backtest_simple.xlsx

Same three 180-day windows and the same trades as export_trades.py, but:
  * NO fees. At 10x leverage a win adds 10% (x1.10), a loss removes 10% (x0.90).
  * Leverage only (no unleveraged % column).
  * Each window starts at $70,000, 10x, full compounding.

Run with:
  uv run --with openpyxl --with pandas --with requests export_trades_simple.py
"""

# --- make repo root + backtest/ subfolders importable (flat layout preserved) ---
import sys as _sys
import pathlib as _pathlib
_ROOT = next(p for p in _pathlib.Path(__file__).resolve().parents
             if (p / "binance_logic.py").exists())
for _d in (_ROOT, _ROOT / "backtest", _ROOT / "backtest" / "runs",
           _ROOT / "backtest" / "export", _ROOT / "backtest" / "checks"):
    if str(_d) not in _sys.path:
        _sys.path.insert(0, str(_d))

from openpyxl import Workbook

from export_trades import (load_windows, style_header, BOLD, HEADER_FILL,
                           WIN_FILL, LOSS_FILL, MONEY_FMT, PRICE_FMT, START_BALANCE)

WIN_MULT = 1.10    # +10% on a win, no fees
LOSS_MULT = 0.90   # -10% on a loss, no fees
_OUTPUT_DIR = _ROOT / "output"
_OUTPUT_DIR.mkdir(exist_ok=True)
XLSX_PATH = str(_OUTPUT_DIR / "mutaz_backtest_simple.xlsx")

PCT_FMT = '0.00"%"'
SIGNED_MONEY = '"$"#,##0.00;"-$"#,##0.00'  # keeps the minus before the $

GEMINI_CLAIM = "Gemini's claim: $70,000 → $472,500 in 6 months."
FEES_NOTE = "Fees are not included. Real results would be lower."

TRADE_COLS = [
    ("Trade #", 9), ("Direction", 11), ("Entry Time (UTC)", 20),
    ("Entry Price", 14), ("Exit Time (UTC)", 20), ("Exit Price", 14),
    ("Result", 9), ("Profit/Loss ($)", 16), ("Balance After", 16),
]

SUMMARY_COLS = [
    ("Period", 22), ("Trades", 9), ("Wins", 7), ("Losses", 8),
    ("Starting Balance", 17), ("Ending Balance", 17),
    ("Profit/Loss ($)", 17), ("Profit/Loss (%)", 16),
]


def write_trade_sheet(ws, trades):
    """One row per trade + a bold total row; returns per-window stats."""
    style_header(ws, TRADE_COLS)

    bal = START_BALANCE
    wins = 0
    for i, t in enumerate(trades, start=1):
        is_win = t["result"] == "TP"
        wins += is_win
        before = bal
        bal *= WIN_MULT if is_win else LOSS_MULT
        pl = bal - before  # +10% of before on a win, -10% on a loss

        values = [
            i, t["direction"], t["entry_time"], t["entry_price"],
            t["exit_time"], t["exit_price"], "WIN" if is_win else "LOSS",
            round(pl, 2), round(bal, 2),
        ]
        row = i + 1
        for c, v in enumerate(values, start=1):
            ws.cell(row=row, column=c, value=v)
        ws.cell(row=row, column=4).number_format = PRICE_FMT
        ws.cell(row=row, column=6).number_format = PRICE_FMT
        ws.cell(row=row, column=8).number_format = SIGNED_MONEY
        ws.cell(row=row, column=9).number_format = MONEY_FMT

        fill = WIN_FILL if is_win else LOSS_FILL
        for c in range(1, len(TRADE_COLS) + 1):
            ws.cell(row=row, column=c).fill = fill

    # ---- total row ----
    n = len(trades)
    losses = n - wins
    ending = bal
    pl_total = ending - START_BALANCE
    pl_pct = (ending / START_BALANCE - 1) * 100

    tr = n + 2  # +1 header, +1 to go below the last trade
    cells = {
        1: "TOTAL",
        2: n,
        3: f"{wins}W / {losses}L",
        4: START_BALANCE,                 # starting balance
        7: f"{pl_pct:+.1f}%",             # profit/loss %
        8: round(pl_total, 2),            # profit/loss $
        9: round(ending, 2),              # ending balance
    }
    for c, v in cells.items():
        cell = ws.cell(row=tr, column=c, value=v)
        cell.font = BOLD
        cell.fill = HEADER_FILL
    ws.cell(row=tr, column=4).number_format = MONEY_FMT
    ws.cell(row=tr, column=8).number_format = SIGNED_MONEY
    ws.cell(row=tr, column=9).number_format = MONEY_FMT

    return {
        "trades": n, "wins": wins, "losses": losses,
        "ending": ending, "pl": pl_total, "pl_pct": pl_pct,
    }


def write_summary(ws, windows, stats_by_window):
    style_header(ws, SUMMARY_COLS)
    for i, w in enumerate(windows, start=1):
        s = stats_by_window[w["name"]]
        row = i + 1
        values = [
            w["name"], s["trades"], s["wins"], s["losses"],
            START_BALANCE, round(s["ending"], 2),
            round(s["pl"], 2), round(s["pl_pct"], 2),
        ]
        for c, v in enumerate(values, start=1):
            ws.cell(row=row, column=c, value=v)
        ws.cell(row=row, column=5).number_format = MONEY_FMT
        ws.cell(row=row, column=6).number_format = MONEY_FMT
        ws.cell(row=row, column=7).number_format = SIGNED_MONEY
        ws.cell(row=row, column=8).number_format = PCT_FMT

    base = len(windows) + 3
    claim = ws.cell(row=base, column=1, value=GEMINI_CLAIM)
    claim.font = BOLD
    ws.cell(row=base + 1, column=1, value=FEES_NOTE)


def main():
    # load_windows returns newest-first; present oldest-first here.
    windows = list(reversed(load_windows()))

    wb = Workbook()
    summary_ws = wb.active
    summary_ws.title = "Summary"

    stats_by_window = {}
    for w in windows:
        ws = wb.create_sheet(title=w["name"])
        stats_by_window[w["name"]] = write_trade_sheet(ws, w["trades"])

    write_summary(summary_ws, windows, stats_by_window)
    wb.save(XLSX_PATH)

    print(f"\nSaved {XLSX_PATH} with sheets: "
          f"{['Summary'] + [w['name'] for w in windows]}\n")
    print(f"{'period':<22}{'W/L':>8}{'ending $':>15}{'P/L %':>9}")
    for w in windows:
        s = stats_by_window[w["name"]]
        wl = f"{s['wins']}/{s['losses']}"
        print(f"{w['name']:<22}{wl:>8}{s['ending']:>15,.0f}{s['pl_pct']:>8.1f}%")


if __name__ == "__main__":
    main()
